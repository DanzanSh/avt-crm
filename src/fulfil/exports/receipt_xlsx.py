"""xlsx для приёмки по плану (Этап 5 плана №3, п.5.2-5.3): шаблон плана для
скачивания, разбор загруженного файла и акт расхождений после завершения приёмки.

Разбор (parse_plan_rows) намеренно не знает про клиента и БД — он просто читает
ячейки и отдаёт сырые строки с номером; сопоставление баркода с товаром клиента и
текст ошибки «строка N — причина» собирает services.receiving.import_plan_xlsx,
у которого есть доступ к каталогу.
"""

import io

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

TEMPLATE_HEADERS = ["Баркод", "Артикул", "Наименование", "Размер", "Цвет", "Ожидается"]
REPORT_HEADERS = ["Баркод", "Наименование", "Размер", "Цвет", "Заявлено", "Принято", "Расхождение", "Статус"]


def _autosize(ws, headers: list[str]) -> None:
    for i, header in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(12, len(header) + 4)


def build_template_xlsx(products: list[dict]) -> bytes:
    """products: [{barcode, vendorCode, name, size, color}] — живые товары клиента."""
    wb = Workbook()
    ws = wb.active
    ws.title = "План приёмки"
    ws.append(TEMPLATE_HEADERS)
    for p in products:
        ws.append(
            [p["barcode"], p.get("vendorCode") or "", p["name"], p.get("size") or "", p.get("color") or "", None]
        )
    _autosize(ws, TEMPLATE_HEADERS)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def parse_plan_rows(file_bytes: bytes) -> list[dict]:
    """Сырые строки листа, кроме заголовка: [{rowNum, barcode, qtyRaw}].
    Полностью пустые строки пропускаются молча — это не ошибка формата."""
    wb = load_workbook(io.BytesIO(file_bytes), data_only=True)
    ws = wb.active
    rows = []
    for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if row is None or all(c is None or str(c).strip() == "" for c in row):
            continue
        barcode = row[0] if len(row) > 0 else None
        qty_raw = row[5] if len(row) > 5 else None
        rows.append(
            {
                "rowNum": idx,
                "barcode": str(barcode).strip() if barcode is not None else "",
                "qtyRaw": qty_raw,
            }
        )
    return rows


def build_discrepancy_report_xlsx(rows: list[dict]) -> bytes:
    """rows: [{barcode, name, size, color, expectedQty, acceptedQty, diff, status}]."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Расхождения"
    ws.append(REPORT_HEADERS)
    for r in rows:
        ws.append(
            [
                r["barcode"], r["name"], r.get("size") or "", r.get("color") or "",
                r["expectedQty"], r["acceptedQty"], r["diff"], r["status"],
            ]
        )
    _autosize(ws, REPORT_HEADERS)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
